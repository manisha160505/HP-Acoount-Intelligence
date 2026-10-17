"""Load the client's user sheet ("Bridge AI users FY26 SNA.xlsx") as users.

The sheet has two tabs:
  Sales  - one row per account: Country, Sales Territory ID, Sales Territory
           Name, EAM, IAM. Every EAM and IAM email named on a row becomes a
           seller with that account; "-" means nobody.
  Admin  - one email per row in column A, with no header row.

Accounts are matched by name: our account names are the client's Sales
Territory Names (scripts/split_account_data.py uploads `name_for_upload`). A
name that does not match exactly is matched only through ALIASES, a list a
person reviewed - never by a guess. Anything else is reported and left out.

Passwords live in a credentials workbook, generated once and then reused: the
production run reads the same file the local run wrote, so the passwords handed
to the client are the ones that work. A password already in the file is never
replaced, and an existing user's password is never touched.
"""

import re
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path

from app.core.security import generate_temporary_password, get_password_hash
from app.services import user_audit

EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")

# Sheet name -> our account name, for the rows whose names differ from ours.
# Checked by hand against the 220 split folders on 17 Oct 2026. Used only when
# the sheet name has no exact match, so a database that does carry the sheet's
# name still matches it directly.
ALIASES = {
    "BUNNINGS GROUP LIMITED- AU": "BUNNINGS GROUP LIMITED",
    "INDONESIA FINANCIAL GROUP - ID": "INDONESIA FINANCIAL GROUP - CATEGORY1 - PARTIAL - ID",
    "TIARA MARGA TRAKINDO - ID": "PT TIARA MARGA TRAKINDO - CATEGORY1 - PARTIAL - ID",
    # Production keeps the pilot account under its original name (17 Oct load:
    # this row was the one unmatched; its seller was given Astra by hand).
    "PT ASTRA INTERNATIONAL TBK - ID": "Astra",
}

ACTOR = {"id": None, "email": "load_sna_users script"}


def norm_name(name) -> str:
    """Case, spacing and the spacing around " - " ignored; nothing else."""
    s = re.sub(r"\s*-\s*", " - ", str(name or ""))
    return re.sub(r"\s+", " ", s).strip().upper()


def norm_email(value) -> str | None:
    s = str(value or "").strip().lower()
    return s if EMAIL_RE.match(s) else None


@dataclass
class SalesRow:
    territory_id: str
    name: str
    emails: list[str]


@dataclass
class Sheet:
    admins: list[str]
    rows: list[SalesRow]
    invalid: list[str] = field(default_factory=list)


def read_sheet(path: str | Path) -> Sheet:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = Sheet(admins=[], rows=[])
    for (value,) in wb["Admin"].iter_rows(max_col=1, values_only=True):
        if value is None or not str(value).strip():
            continue
        email = norm_email(value)
        if email:
            if email not in sheet.admins:
                sheet.admins.append(email)
        else:
            sheet.invalid.append("Admin: %r" % value)
    for r in wb["Sales"].iter_rows(min_row=2, values_only=True):
        if not r or not any(r):
            continue
        name = str(r[2] or "").strip()
        emails = []
        for raw in (r[3], r[4]):
            if raw is None or str(raw).strip() in ("", "-"):
                continue
            email = norm_email(raw)
            if email:
                emails.append(email)
            else:
                sheet.invalid.append("%s: %r" % (name, raw))
        sheet.rows.append(SalesRow(str(r[1] or ""), name, emails))
    return sheet


@dataclass
class Person:
    email: str
    role: str
    account_ids: list[str] = field(default_factory=list)
    account_names: list[str] = field(default_factory=list)
    status: str = "new"            # new | exists | role_differs
    current_role: str | None = None


@dataclass
class Plan:
    people: list[Person]
    matched: list[tuple[str, str, str]]          # (sheet name, account name, exact|alias)
    unmatched: list[str]                         # sheet names with no account
    accounts_not_in_sheet: list[str]
    notes: list[str] = field(default_factory=list)
    account_names: dict[str, str] = field(default_factory=dict)   # id -> name


def build_plan(sheet: Sheet, accounts: list[dict], existing: dict[str, dict]) -> Plan:
    """`accounts`: [{"id", "name"}] as in the database; `existing`: email -> user."""
    by_name = {norm_name(a["name"]): a for a in accounts}
    matched, unmatched, used = [], [], set()
    seller_accounts: dict[str, list[dict]] = {}
    for row in sheet.rows:
        acct, how = by_name.get(norm_name(row.name)), "exact"
        if acct is None and row.name.strip() in ALIASES:
            acct, how = by_name.get(norm_name(ALIASES[row.name.strip()])), "alias"
        if acct is None:
            unmatched.append(row.name)
            continue
        matched.append((row.name, acct["name"], how))
        used.add(acct["id"])
        for email in row.emails:
            lst = seller_accounts.setdefault(email, [])
            if acct not in lst:
                lst.append(acct)
    # Emails only on unmatched rows still get an account, with nothing assigned,
    # so they appear under "No accounts" on the Users page instead of vanishing.
    for row in sheet.rows:
        for email in row.emails:
            seller_accounts.setdefault(email, [])

    notes = []
    people = []
    for email in sheet.admins:
        if email in seller_accounts:
            notes.append("%s is on both tabs; made an admin (admins see every account)." % email)
        people.append(Person(email, "admin"))
    for email, accts in sorted(seller_accounts.items()):
        if email in sheet.admins:
            continue
        accts = sorted(accts, key=lambda a: a["name"])
        people.append(Person(email, "user", [a["id"] for a in accts], [a["name"] for a in accts]))
    for p in people:
        user = existing.get(p.email)
        if user:
            p.current_role = user.get("role") or "user"
            p.status = "exists" if p.current_role == p.role else "role_differs"
    return Plan(people, matched, unmatched,
                sorted(a["name"] for a in accounts if a["id"] not in used), notes,
                {a["id"]: a["name"] for a in accounts})


# -------------------------------------------------------------- credentials

CRED_SHEET = "Credentials"
CRED_HEADER = ["Email", "Role", "Temporary password", "Accounts", "Account names"]


def load_credentials(path: str | Path) -> dict[str, str]:
    p = Path(path)
    if not p.exists():
        return {}
    import openpyxl
    ws = openpyxl.load_workbook(p, read_only=True, data_only=True)[CRED_SHEET]
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and row[0] and row[2]:
            out[str(row[0]).strip().lower()] = str(row[2])
    return out


def fill_passwords(plan: Plan, passwords: dict[str, str]) -> list[str]:
    """Give every new person without one a generated password. Returns the
    emails that got one. Never replaces a password already in `passwords`."""
    added = []
    for p in plan.people:
        if p.status == "new" and p.email not in passwords:
            passwords[p.email] = generate_temporary_password()
            added.append(p.email)
    return added


def write_credentials(path: str | Path, plan: Plan, passwords: dict[str, str]) -> None:
    """Every email with a password, the ones no longer in the plan included: a
    password handed out once stays on record."""
    import openpyxl
    from openpyxl.styles import Font
    people = {p.email: p for p in plan.people}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = CRED_SHEET
    ws.append(CRED_HEADER)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    order = sorted(passwords, key=lambda e: (people[e].role != "admin" if e in people else 2, e))
    for email in order:
        p = people.get(email)
        ws.append([email, p.role if p else "(not in sheet)", passwords[email],
                   len(p.account_ids) if p else "", ", ".join(p.account_names) if p else ""])
    for col, width in zip("ABCDE", (38, 10, 20, 10, 90)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    if plan.unmatched:
        um = wb.create_sheet("Unmatched accounts")
        um.append(["Sales Territory Name in the sheet with no matching account"])
        for name in plan.unmatched:
            um.append([name])
    wb.save(path)


# -------------------------------------------------------------------- apply

def apply(db, plan: Plan, passwords: dict[str, str], *, super_admin_email: str,
          super_admin_password: str | None, sync_accounts: bool = False) -> dict[str, int]:
    """Create what the plan says is new. Existing users keep their password and
    role; with `sync_accounts`, existing sellers get the sheet's accounts."""
    missing = [p.email for p in plan.people if p.status == "new" and p.email not in passwords]
    if missing:
        raise ValueError("no password in the credentials file for: %s" % ", ".join(missing))
    now = datetime.now(UTC)
    counts = {"created": 0, "accounts_synced": 0, "super_admin_created": 0}

    sa = db["users"].find_one({"email": super_admin_email})
    if sa is None:
        if not super_admin_password:
            raise ValueError("SUPER_ADMIN_PASSWORD is not set and %s does not exist yet"
                             % super_admin_email)
        doc = {"email": super_admin_email, "full_name": super_admin_email, "role": "super_admin",
               "password_hash": get_password_hash(super_admin_password), "account_ids": [],
               "is_active": True, "must_change_password": False, "created_at": now,
               "created_by": "load_sna_users", "last_login_at": None}
        doc["_id"] = db["users"].insert_one(doc).inserted_id
        user_audit.record(db, ACTOR, user_audit.USER_CREATED, doc, {"role": "super_admin"})
        counts["super_admin_created"] = 1
    elif sa.get("role") != "super_admin":
        raise ValueError("%s exists with role %r; refusing to promote it silently"
                         % (super_admin_email, sa.get("role")))

    for p in plan.people:
        if p.status == "new":
            doc = {"email": p.email, "full_name": p.email, "role": p.role,
                   "password_hash": get_password_hash(passwords[p.email]),
                   "account_ids": p.account_ids, "is_active": True,
                   "must_change_password": True, "created_at": now,
                   "created_by": "load_sna_users", "last_login_at": None}
            doc["_id"] = db["users"].insert_one(doc).inserted_id
            user_audit.record(db, ACTOR, user_audit.USER_CREATED, doc,
                              {"role": p.role, "accounts": p.account_names})
            counts["created"] += 1
        elif sync_accounts and p.status == "exists" and p.role == "user":
            user = db["users"].find_one({"email": p.email})
            old = [str(a) for a in user.get("account_ids") or []]
            if old != p.account_ids:
                db["users"].update_one({"_id": user["_id"]}, {"$set": {
                    "account_ids": p.account_ids, "updated_at": now}})
                name = plan.account_names.get
                user_audit.record(db, ACTOR, user_audit.ACCOUNTS_CHANGED, user, {
                    "added": [name(i, i) for i in p.account_ids if i not in old],
                    "removed": [name(i, i) for i in old if i not in p.account_ids],
                })
                counts["accounts_synced"] += 1
    return counts
